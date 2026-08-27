"""
Tests for the OCR extraction post-pass (core/text_extraction.py).

The DocAble HTTP service is faked throughout — these tests prove the
fork-side logic: hash caching, verdict->column mapping, and Excel updates.
"""
import json

import openpyxl
import pytest

from core.text_extraction import (
    ExtractionPass,
    requires_ocr_cell_value,
    update_excel_requires_ocr,
)


class FakeClient:
    """Stands in for the HTTP client; counts calls to prove caching."""

    def __init__(self, verdict="scanned", requires_ocr=True):
        self.calls = 0
        self._body = {
            "verdict": {
                "requires_ocr": requires_ocr,
                "classification": verdict,
                "scanned_pages": [1] if requires_ocr else [],
                "scanned_page_count": 1 if requires_ocr else 0,
            },
            "page_count": 1,
            "report_url": "/audits/x-abc123/report.html",
        }

    def extract(self, path):
        self.calls += 1
        return self._body


@pytest.fixture()
def course(tmp_path):
    """A minimal downloaded-course folder: one PDF + the JSON export."""
    pdf = tmp_path / "docs" / "syllabus.pdf"
    pdf.parent.mkdir()
    pdf.write_bytes(b"%PDF-1.4 fake content for hashing")
    export = {
        "course_id": "123",
        "content": {"documents": {"documents": [
            {"title": "Syllabus", "url": "https://x/files/1", "save_path": str(pdf)},
            {"title": "Missing", "url": "https://x/files/2", "save_path": str(tmp_path / "gone.pdf")},
        ], "document_sites": []}},
    }
    (tmp_path / "123.json").write_text(json.dumps(export))
    return tmp_path


class TestExtractionPass:
    def test_extracts_each_existing_pdf_once(self, course):
        client = FakeClient()
        results = ExtractionPass(client, course).run()
        assert client.calls == 1
        assert results["https://x/files/1"]["classification"] == "scanned"
        assert results["https://x/files/1"]["report_url"] == "/audits/x-abc123/report.html"

    def test_missing_files_reported_not_crashed(self, course):
        results = ExtractionPass(FakeClient(), course).run()
        assert results["https://x/files/2"]["error"] == "file_missing"

    def test_second_run_hits_cache_not_service(self, course):
        client = FakeClient()
        ExtractionPass(client, course).run()
        ExtractionPass(client, course).run()
        assert client.calls == 1  # cache keyed by content hash

    def test_changed_file_re_extracts(self, course):
        client = FakeClient()
        ExtractionPass(client, course).run()
        (course / "docs" / "syllabus.pdf").write_bytes(b"%PDF-1.4 remediated!")
        ExtractionPass(client, course).run()
        assert client.calls == 2

    def test_service_error_recorded_and_not_cached(self, course):
        class Boom:
            calls = 0
            def extract(self, path):
                Boom.calls += 1
                raise ConnectionError("service down")
        boom = Boom()
        results = ExtractionPass(boom, course).run()
        assert results["https://x/files/1"]["error"] == "service_unreachable"
        ExtractionPass(boom, course).run()
        assert Boom.calls == 2  # failure was not cached as a verdict


class TestCellValue:
    def test_requires_ocr_maps_to_yes(self):
        assert requires_ocr_cell_value({"requires_ocr": True, "classification": "scanned"}) == "Yes"

    def test_digital_maps_to_no(self):
        assert requires_ocr_cell_value({"requires_ocr": False, "classification": "digital"}) == "No"

    def test_no_text_maps_to_no(self):
        assert requires_ocr_cell_value({"requires_ocr": False, "classification": "no_text"}) == "No"

    def test_error_maps_to_none_leave_cell_alone(self):
        assert requires_ocr_cell_value({"error": "service_unreachable"}) is None


class TestExcelUpdate:
    def _workbook(self, path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Documents"
        ws.append(["Title", "Url", "Requires OCR", "Notes"])
        ws.append(["Syllabus", "https://x/files/1", "No", None])
        ws.append(["Slides", "https://x/files/9", "No", None])
        wb.save(path)
        return path

    def test_updates_matching_rows_only(self, tmp_path):
        xl = self._workbook(tmp_path / "c.xlsx")
        n = update_excel_requires_ocr(
            xl, {"https://x/files/1": {"requires_ocr": True, "classification": "scanned"}}
        )
        assert n == 1
        ws = openpyxl.load_workbook(xl)["Documents"]
        assert ws.cell(row=2, column=3).value == "Yes"
        assert ws.cell(row=3, column=3).value == "No"   # untouched

    def test_error_verdicts_leave_cells_untouched(self, tmp_path):
        xl = self._workbook(tmp_path / "c.xlsx")
        n = update_excel_requires_ocr(xl, {"https://x/files/1": {"error": "service_unreachable"}})
        assert n == 0
        ws = openpyxl.load_workbook(xl)["Documents"]
        assert ws.cell(row=2, column=3).value == "No"


class TestOfficeRouting:
    def test_docx_is_sent_to_the_service(self, course, tmp_path):
        import json
        docx = course / "docs" / "notes.docx"
        docx.write_bytes(b"PK\x03\x04 fake docx payload")
        export = json.loads((course / "123.json").read_text())
        export["content"]["documents"]["documents"].append(
            {"title": "Notes", "url": "https://x/files/3", "save_path": str(docx)}
        )
        (course / "123.json").write_text(json.dumps(export))
        client = FakeClient()
        results = ExtractionPass(client, course).run()
        assert "https://x/files/3" in results
        assert client.calls == 2  # pdf + docx


class TestMarkdownPersistence:
    def test_extracted_markdown_saved_per_document(self, course):
        class MdClient(FakeClient):
            def extract(self, path):
                body = super().extract(path)
                return {**body, "markdown": "## Redox Titrations\n\nDilute the solution."}
        results = ExtractionPass(MdClient(), course).run()
        md_rel = results["https://x/files/1"]["markdown_file"]
        md_path = course / ".manifest" / md_rel
        assert md_path.is_file()
        assert "## Redox Titrations" in md_path.read_text()
