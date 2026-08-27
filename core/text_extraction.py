"""
Text Extraction Post-Pass
=========================

Connects a completed course scan to the doc-extract OCR service and turns
its per-document verdicts into the Excel report's "Requires OCR" column —
a judgment that column has always asked a human to make by eye.

Architecture
------------
This is a POST-download pass, run against a course folder that already
holds downloaded files plus the JSON export::

    {course_folder}/
    ├── 123456.json          <- save_content_as_json() export (has save_path per doc)
    ├── 123456.xlsx          <- save_as_excel() report (has Requires OCR column)
    ├── .manifest/
    │   ├── extraction_cache.json    <- verdicts keyed by file sha256
    │   └── extraction_results.json  <- last run, keyed by canvas URL
    └── {date}/.../file.pdf

Running it after the scan (rather than inline during download) keeps scan
time predictable, works when the OCR service is down (errors are recorded,
never cached), and makes re-runs cheap: verdicts are cached by content
hash, so only new or remediated files hit the GPU again.

The service is a separate self-hosted project (doc-extract) reached over
HTTP; nothing in this module imports torch or touches a model.

See Also
--------
- tools.extract_pass : CLI wrapper for this module
- tools.export_to_excel : Defines the tracking columns this pass fills in
"""

from __future__ import annotations

import hashlib
import json
import logging
import os
from glob import glob

import openpyxl
import requests

log = logging.getLogger(__name__)

REQUEST_TIMEOUT_SECONDS = 600  # a large scanned PDF is minutes of GPU time
CACHE_FILE = "extraction_cache.json"
RESULTS_FILE = "extraction_results.json"

# Formats the doc-extract service accepts: PDFs and images natively, office
# and OpenDocument formats via its LibreOffice conversion path.
EXTRACTABLE_SUFFIXES = (
    ".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff",
    ".docx", ".pptx", ".xlsx", ".doc", ".ppt", ".xls",
    ".odt", ".odp", ".ods", ".rtf",
)


class ExtractionClient:
    """Thin HTTP client for the doc-extract service."""

    def __init__(self, endpoint: str, session: requests.Session | None = None):
        self.endpoint = endpoint.rstrip("/")
        self.session = session or requests.Session()

    def extract(self, file_path: str) -> dict:
        with open(file_path, "rb") as f:
            response = self.session.post(
                f"{self.endpoint}/extract",
                files={"file": (os.path.basename(file_path), f)},
                timeout=REQUEST_TIMEOUT_SECONDS,
            )
        response.raise_for_status()
        return response.json()


def _sha256(path: str) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _load_documents(course_folder: str) -> list[dict]:
    """Find the course JSON export and return its document entries."""
    for candidate in sorted(glob(os.path.join(course_folder, "*.json"))):
        try:
            with open(candidate, encoding="utf-8") as f:
                data = json.load(f)
        except (OSError, ValueError):
            continue
        documents = (
            data.get("content", {}).get("documents", {}).get("documents")
            if isinstance(data, dict) else None
        )
        if documents is not None:
            return documents
    raise FileNotFoundError(
        f"No course JSON export found in {course_folder} — run the scan with "
        f"JSON output first (the export provides the URL<->file mapping)."
    )


class ExtractionPass:
    """
    One post-pass over a course folder: extract, cache, report.

    Results are keyed by the document's Canvas URL — the join key shared
    with the Excel report, which carries URLs but not local paths.
    """

    def __init__(self, client, course_folder):
        self.client = client
        self.course_folder = str(course_folder)
        self._manifest_dir = os.path.join(self.course_folder, ".manifest")
        self._cache_path = os.path.join(self._manifest_dir, CACHE_FILE)

    def run(self) -> dict:
        cache = self._read_json(self._cache_path)
        results = {}

        for doc in _load_documents(self.course_folder):
            url, save_path = doc.get("url"), doc.get("save_path")
            if not url or not save_path:
                continue
            if not save_path.lower().endswith(EXTRACTABLE_SUFFIXES):
                continue
            if not os.path.isfile(save_path):
                results[url] = {"error": "file_missing", "save_path": save_path}
                continue

            file_hash = _sha256(save_path)
            if file_hash in cache:
                results[url] = cache[file_hash]
                continue

            try:
                body = self.client.extract(save_path)
                verdict = dict(body["verdict"])
                if body.get("report_url"):
                    verdict["report_url"] = body["report_url"]
                if body.get("accessibility"):
                    verdict["score"] = body["accessibility"]["score"]
                    verdict["band"] = body["accessibility"]["band"]
            except Exception as exc:  # noqa: BLE001 - any failure means "no verdict"
                log.warning(f"Extraction failed for {save_path}: {exc}")
                results[url] = {"error": "service_unreachable", "detail": str(exc)}
                continue  # deliberately NOT cached: retry on next run

            # Persist the recovered text: for a scanned document this is the
            # remediation starting point, and downstream projections (Obsidian,
            # web UI) render it. Stored by hash so re-runs and renames dedupe.
            markdown = body.get("markdown", "")
            if markdown:
                md_dir = os.path.join(self._manifest_dir, "extracted_md")
                os.makedirs(md_dir, exist_ok=True)
                with open(os.path.join(md_dir, f"{file_hash}.md"), "w", encoding="utf-8") as f:
                    f.write(markdown)
                verdict["markdown_file"] = os.path.join("extracted_md", f"{file_hash}.md")

            verdict["sha256"] = file_hash
            cache[file_hash] = verdict
            results[url] = verdict

        self._write_json(self._cache_path, cache)
        self._write_json(os.path.join(self._manifest_dir, RESULTS_FILE), results)
        return results

    @staticmethod
    def _read_json(path: str) -> dict:
        try:
            with open(path, encoding="utf-8") as f:
                return json.load(f)
        except (OSError, ValueError):
            return {}

    def _write_json(self, path: str, data: dict) -> None:
        os.makedirs(self._manifest_dir, exist_ok=True)
        with open(path, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2)


def requires_ocr_cell_value(verdict: dict) -> str | None:
    """
    Map a document verdict to the Excel dropdown's vocabulary.

    Errors map to None — the pass must never overwrite a human's earlier
    judgment with a guess just because the service was unreachable.
    """
    if "error" in verdict:
        return None
    return "Yes" if verdict.get("requires_ocr") else "No"


def update_excel_requires_ocr(excel_path, url_to_verdict: dict) -> int:
    """
    Fill the Documents sheet's "Requires OCR" column in place.

    Rows are matched by the Url column. Returns how many cells were set.
    """
    wb = openpyxl.load_workbook(excel_path, keep_vba=str(excel_path).endswith(".xlsm"))
    sheet = wb["Documents"]
    headers = {cell.value: cell.column for cell in sheet[1] if cell.value}
    url_col, ocr_col = headers.get("Url"), headers.get("Requires OCR")
    if not url_col or not ocr_col:
        raise ValueError("Documents sheet is missing the Url or Requires OCR column")

    updated = 0
    for row in range(2, sheet.max_row + 1):
        verdict = url_to_verdict.get(sheet.cell(row=row, column=url_col).value)
        if not verdict:
            continue
        value = requires_ocr_cell_value(verdict)
        if value is None:
            continue
        sheet.cell(row=row, column=ocr_col).value = value
        updated += 1

    wb.save(excel_path)
    return updated
